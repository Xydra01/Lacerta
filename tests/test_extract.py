"""Document extract helpers (V1.55)."""

from __future__ import annotations

from pathlib import Path

import pytest

from lacerta.storage.extract import (
    ExtractError,
    extract_text,
    is_extractable,
    missing_optional_deps,
    normalized_source_name,
)


def test_extract_plain_and_html(tmp_path: Path) -> None:
    md = tmp_path / "a.md"
    md.write_text("# Hello\n\nworld planted_token_md\n", encoding="utf-8")
    assert "planted_token_md" in extract_text(md)

    html = tmp_path / "b.html"
    html.write_text(
        "<html><body><h1>Title</h1><p>planted_token_html</p>"
        "<script>ignore()</script></body></html>",
        encoding="utf-8",
    )
    text = extract_text(html)
    assert "planted_token_html" in text
    assert "ignore" not in text


def test_unsupported_and_legacy_doc(tmp_path: Path) -> None:
    bad = tmp_path / "x.pptx"
    bad.write_bytes(b"nope")
    with pytest.raises(ExtractError) as ei:
        extract_text(bad)
    assert ei.value.code == "unsupported_source"

    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"legacy")
    with pytest.raises(ExtractError) as ej:
        extract_text(legacy)
    assert ej.value.code == "unsupported_source"
    assert "docx" in ej.value.message.lower()


def test_normalized_source_name() -> None:
    assert normalized_source_name(Path("Book.PDF")) == "Book.md"
    assert is_extractable(Path("a.pdf"))


@pytest.mark.skipif(
    missing_optional_deps(Path("x.pdf")) is not None,
    reason="pypdf not installed",
)
def test_extract_pdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "planted.pdf"
    pdf_path.write_bytes(
        b"""%PDF-1.4
1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj
2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj
3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]
/Contents 4 0 R /Resources<< /Font<< /F1 5 0 R >> >> >>endobj
4 0 obj<< /Length 58 >>stream
BT /F1 24 Tf 72 720 Td (LACERTA_PDF_PLANTED_FACT) Tj ET
endstream
endobj
5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000266 00000 n 
0000000373 00000 n 
trailer<< /Size 6 /Root 1 0 R >>
startxref
446
%%EOF
"""
    )
    text = extract_text(pdf_path)
    assert "LACERTA_PDF_PLANTED_FACT" in text.replace(" ", "")


@pytest.mark.skipif(
    missing_optional_deps(Path("x.docx")) is not None,
    reason="python-docx not installed",
)
def test_extract_docx(tmp_path: Path) -> None:
    import docx

    path = tmp_path / "note.docx"
    doc = docx.Document()
    doc.add_heading("Doc Title", level=1)
    doc.add_paragraph("LACERTA_DOCX_PLANTED_FACT appears here.")
    doc.save(path)
    text = extract_text(path)
    assert "LACERTA_DOCX_PLANTED_FACT" in text


def test_corpus_extract_html_to_sources(tmp_path: Path) -> None:
    from lacerta.core.capabilities import CapabilityContext, run_capability
    from lacerta.storage import corpus as corpus_storage
    from lacerta.workers.corpus import capabilities as _cc  # noqa: F401

    html = tmp_path / "lecture.html"
    html.write_text(
        "<html><body><p>LACERTA_HTML_CORPUS_FACT in a lecture.</p></body></html>",
        encoding="utf-8",
    )
    ctx = CapabilityContext(
        instance_id="default",
        course_id="html-course",
        data_root=str(tmp_path),
        attachments=[str(html)],
        extra={"data_root": str(tmp_path)},
    )
    result = run_capability("corpus.extract", ctx, {})
    assert result.ok, result.error_message
    root = corpus_storage.resolve_course_corpus_root(tmp_path, "default", "html-course")
    written = list(corpus_storage.sources_dir(root).glob("*.md"))
    assert written
    assert "LACERTA_HTML_CORPUS_FACT" in written[0].read_text(encoding="utf-8")
